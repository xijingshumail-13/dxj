"""Regression checks use synthetic names and temporary workbooks only."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import lottery


class CliEncodingTests(unittest.TestCase):
    def run_cli(self, *arguments):
        environment = os.environ.copy()
        environment.update(PYTHONIOENCODING="cp1252:strict", PYTHONUTF8="0")
        with tempfile.TemporaryDirectory(prefix="lottery-encoding-") as directory:
            return subprocess.run(
                [sys.executable, str(Path(lottery.__file__).resolve()), *arguments],
                cwd=directory, env=environment, capture_output=True, timeout=15,
            )

    def test_help_remains_readable_when_stdout_starts_as_cp1252(self):
        result = self.run_cli("--help")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("大学节 SVG 抽签", result.stdout.decode("utf-8"))
        self.assertEqual(result.stderr, b"")

    def test_missing_roster_error_remains_readable_with_cp1252(self):
        result = self.run_cli("--input", "不存在的测试名单.xlsx", "--no-browser")
        self.assertEqual(result.returncode, 1, result.stderr)
        message = result.stderr.decode("utf-8")
        self.assertIn("无法启动：找不到名单", message)
        self.assertIn("不存在的测试名单.xlsx", message)
        self.assertNotIn("UnicodeEncodeError", message)

    def test_argparse_error_preserves_chinese_argument_with_cp1252(self):
        result = self.run_cli("--port", "不是端口")
        self.assertEqual(result.returncode, 2, result.stderr)
        message = result.stderr.decode("utf-8")
        self.assertIn("不是端口", message)
        self.assertNotIn("UnicodeEncodeError", message)


class LotteryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.directory = Path(self.tmp.name)
        self.output = self.directory / "result.xlsx"

    def workbook(self, rows, name="student.xlsx"):
        path = self.directory / name
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        wb.close()
        return path

    def rows(self, path=None):
        wb = load_workbook(path or self.output, read_only=True, data_only=False)
        try:
            return list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()

    def test_reader_without_header_retains_first_student(self):
        path = self.workbook([["  同学甲  "], [None], ["同学乙"]])
        self.assertEqual(lottery.read_students(path), ["同学甲", "同学乙"])

    def test_reader_blank_first_row_and_recognized_header(self):
        for rows in ([[None], ["同学甲"], ["同学乙"]],
                     [[None], ["学生姓名"], ["同学甲"], ["同学乙"]]):
            with self.subTest(rows=rows):
                self.assertEqual(lottery.read_students(self.workbook(rows)), ["同学甲", "同学乙"])

    def test_reader_name_column_and_duplicate_names(self):
        path = self.workbook([["序号", "姓名"], [1, "重名同学"], [2, "重名同学"], [3, "同学丙"]])
        self.assertEqual(lottery.read_students(path), ["重名同学", "重名同学", "同学丙"])

    def test_reader_rejects_empty_workbook(self):
        with self.assertRaisesRegex(ValueError, "名单为空"):
            lottery.read_students(self.workbook([[None]]))

    def test_random_plans_preserve_round_and_pair_rules(self):
        for count in (1, 2, 24):
            for trial in range(100):
                with self.subTest(count=count, trial=trial):
                    plan = lottery.make_plan(count, lottery.UNIVERSITIES)
                    self.assertEqual(len(plan), count)
                    self.assertTrue(all(len(pair) == 2 and pair[0] != pair[1] for pair in plan))
                    for round_index in (0, 1):
                        column = [pair[round_index] for pair in plan]
                        self.assertEqual(len(set(column)), count)
                        self.assertTrue(set(column) <= set(lottery.UNIVERSITIES))

    def test_all_48_draws_complete_with_no_repetition_within_either_round(self):
        names = [f"测试同学{i + 1}" for i in range(24)]
        game = lottery.Lottery(names, self.output)
        for cursor in range(48):
            result = game.draw(game.session_id, cursor)
            self.assertEqual(result["state"]["cursor"], cursor + 1)
            self.assertEqual(result["draw"]["student_index"], cursor // 2)
            self.assertEqual(result["draw"]["round_index"], cursor % 2)
        state = game.snapshot()
        self.assertTrue(state["done"])
        self.assertEqual(state["completed_students"], 24)
        self.assertEqual(state["available_universities"], [])
        for round_index in (0, 1):
            self.assertEqual({p[round_index] for p in state["results"]}, set(lottery.UNIVERSITIES))
        self.assertTrue(all(a != b for a, b in state["results"]))
        self.assertEqual(self.rows()[1:], [(i + 1, name, *state["results"][i]) for i, name in enumerate(names)])
        with self.assertRaises(lottery.ConflictError):
            game.draw(game.session_id, 48)

    def test_resume_after_first_draw_preserves_unrevealed_plan(self):
        game = lottery.Lottery(["同学甲", "同学乙"], self.output)
        plan = deepcopy(game.plan)
        game.draw(game.session_id, 0)
        restarted = lottery.Lottery(game.students, self.output)
        self.assertTrue(restarted.restored)
        self.assertEqual(restarted.cursor, 1)
        self.assertEqual(restarted.plan, plan)
        self.assertEqual(restarted.snapshot()["results"], [[plan[0][0], None], [None, None]])
        self.assertNotEqual(restarted.session_id, game.session_id)
        second = restarted.draw(restarted.session_id, 1)
        self.assertEqual(second["draw"]["university"], plan[0][1])

    def test_completed_file_restores_as_complete(self):
        game = lottery.Lottery(["同学甲"], self.output)
        for cursor in range(2):
            game.draw(game.session_id, cursor)
        restarted = lottery.Lottery(game.students, self.output)
        self.assertTrue(restarted.snapshot()["done"])
        self.assertEqual(restarted.plan, game.plan)
        with self.assertRaises(lottery.ConflictError):
            restarted.draw(restarted.session_id, 2)

    def test_failed_save_preserves_previous_file_progress_and_retry(self):
        game = lottery.Lottery(["同学甲", "同学乙"], self.output)
        game.draw(game.session_id, 0)
        previous = self.output.read_bytes()
        with patch("lottery.os.replace", side_effect=PermissionError("file is open")):
            with self.assertRaises(PermissionError):
                game.draw(game.session_id, 1)
        self.assertEqual(game.cursor, 1)
        self.assertEqual(self.output.read_bytes(), previous)
        self.assertEqual(self.rows()[1][3], None)
        self.assertEqual(list(self.directory.glob(".lottery-*.xlsx")), [])
        retry = game.draw(game.session_id, 1)
        self.assertEqual(retry["state"]["cursor"], 2)
        self.assertEqual(self.rows()[1][3], game.plan[0][1])

    def test_concurrent_repeated_cursor_commits_only_one_draw(self):
        game = lottery.Lottery(["同学甲", "同学乙"], self.output)
        gate = threading.Barrier(2)

        def attempt():
            gate.wait(timeout=5)
            try:
                game.draw(game.session_id, 0)
                return "saved"
            except lottery.ConflictError:
                return "conflict"

        with ThreadPoolExecutor(max_workers=2) as executor:
            outcomes = list(executor.map(lambda _: attempt(), range(2)))
        self.assertCountEqual(outcomes, ["saved", "conflict"])
        self.assertEqual(game.cursor, 1)
        self.assertEqual(sum(value is not None for row in self.rows()[1:] for value in row[2:]), 1)

    def test_mismatched_roster_is_refused_without_overwrite(self):
        game = lottery.Lottery(["同学甲", "同学乙"], self.output)
        game.draw(game.session_id, 0)
        previous = self.output.read_bytes()
        with self.assertRaisesRegex(ValueError, "学生名单或姓名顺序已变化"):
            lottery.Lottery(["同学乙", "同学甲"], self.output)
        self.assertEqual(self.output.read_bytes(), previous)

    def test_formula_like_student_name_is_exported_as_literal_text(self):
        game = lottery.Lottery(["=1+1"], self.output)
        game.draw(game.session_id, 0)
        wb = load_workbook(self.output, data_only=False)
        try:
            self.assertEqual(wb.active["B2"].value, "=1+1")
            self.assertEqual(wb.active["B2"].data_type, "s")
        finally:
            wb.close()
        self.assertEqual(lottery.Lottery(game.students, self.output).cursor, 1)

    def test_overcapacity_reports_problem_and_never_writes_results(self):
        game = lottery.Lottery([f"测试同学{i}" for i in range(25)], self.output)
        state = game.snapshot()
        self.assertIn("名额不足", state["setup_error"])
        self.assertEqual(state["cursor"], 0)
        self.assertFalse(state["done"])
        self.assertEqual(state["results"], [[None, None]] * 25)
        with self.assertRaisesRegex(ValueError, "名额不足"):
            game.draw(game.session_id, 0)
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
