"""大学节 SVG 抽签。安装依赖后运行 python lottery.py。"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import secrets
import sys
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlsplit
import webbrowser

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    raise SystemExit("缺少 openpyxl，请先运行：python -m pip install -r requirements.txt")

# PyInstaller 将网页解压到 _MEIPASS，名单与结果必须放在 exe 旁边。
BASE_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
ASSET_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
INPUT_FILE = BASE_DIR / "student.xlsx"
OUTPUT_FILE = BASE_DIR / "result.xlsx"
UNIVERSITIES = [
    "北京大学", "清华大学", "复旦大学", "上海交通大学", "南京大学", "浙江大学",
    "中国科学技术大学", "哈尔滨工业大学（本部）", "西安交通大学", "中国科学院大学",
    "中国人民大学", "北京航空航天大学", "北京理工大学", "北京师范大学", "南开大学",
    "同济大学", "东南大学", "武汉大学", "华中科技大学", "厦门大学", "天津大学",
    "大连理工大学", "华南理工大学", "四川大学",
]
HEADERS = ["序号", "学生姓名", "上半场", "下半场"]
NAME_HEADERS = {"姓名", "学生姓名", "学生", "名字", "name", "student", "student name"}
METADATA_PREFIX = "university-lottery-v2\n"
RULE_TEXT = "每人上、下半场各抽一所不同大学；同一半场内，高校不重复。"


def read_students(path: Path = INPUT_FILE) -> list[str]:
    """读取首张表；识别前十行的姓名表头，否则取 A 列，保留重名。"""
    path = Path(path)
    if not path.is_file():
        raise ValueError(f"找不到名单：{path}。请把 student.xlsx 放在程序所在文件夹。")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 {path.name}，请确认它是有效的 Excel 文件。") from exc
    try:
        ws = wb.worksheets[0]
        name_col, start_row = 1, 1
        found_header = False
        for row_number, row in enumerate(ws.iter_rows(max_row=10, max_col=30, values_only=True), 1):
            for col_number, value in enumerate(row, 1):
                if isinstance(value, str) and value.strip().lower() in NAME_HEADERS:
                    name_col, start_row, found_header = col_number, row_number + 1, True
                    break
            if found_header:
                break
        students = []
        for row in ws.iter_rows(min_row=start_row, min_col=name_col,
                                max_col=name_col, values_only=True):
            value = row[0]
            if value is not None and str(value).strip():
                if not isinstance(value, str):
                    raise ValueError("姓名列包含数字或日期；请使用“姓名”表头指明姓名列。")
                students.append(value.strip())
        if not students:
            raise ValueError("名单为空。请在第一张工作表的 A 列填写姓名，或使用“姓名”表头。")
        return students
    finally:
        wb.close()


def fingerprint(value) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False).encode("utf-8")).hexdigest()


def make_plan(count: int, universities: list[str]) -> list[list[str]]:
    """每场不放回抽样，第二场拒绝任何同学再次抽到原大学的方案。

    先确定完整方案，避免最后一位只剩自己上半场的大学可抽。
    浏览器每次只收到已揭晓的结果，不接收后续方案。
    """
    if count < 1 or count > len(universities) or len(universities) < 2:
        raise ValueError("每位同学需要两所不同大学，且每场大学数量不能少于学生人数。")
    if len(set(universities)) != len(universities):
        raise ValueError("大学列表中存在重复名称，请先删除重复项。")
    rng = secrets.SystemRandom()
    first = rng.sample(universities, count)
    while True:
        second = rng.sample(universities, count)
        if all(a != b for a, b in zip(first, second)):
            return [list(pair) for pair in zip(first, second)]


class ConflictError(ValueError):
    """旧页面或重复点击不能产生额外抽签。"""


class Lottery:
    def __init__(self, students: list[str], output_path: Path = OUTPUT_FILE,
                 universities: list[str] | None = None):
        self.students = list(students)
        self.universities = list(UNIVERSITIES if universities is None else universities)
        self.output_path = Path(output_path)
        self.session_id = secrets.token_urlsafe(24)
        self.lock = threading.RLock()
        self.cursor = 0
        self.restored = False
        self.plan: list[list[str]] = []
        self.setup_error = ""
        if not self.students:
            self.setup_error = "名单为空，请填写 student.xlsx 后重新启动程序。"
        elif len(set(self.universities)) != len(self.universities):
            self.setup_error = "大学列表包含重复名称，请修改 UNIVERSITIES 后重新启动。"
        elif len(self.universities) < 2:
            self.setup_error = "至少需要两所不同大学，才能安排上、下半场。"
        elif len(self.students) > len(self.universities):
            self.setup_error = (
                f"当前有 {len(self.students)} 位同学、{len(self.universities)} 所大学。"
                "同一半场每所大学只安排一人，名额不足。"
                "请调整 student.xlsx 或扩充程序中的大学列表，然后重新启动。"
            )
        if self.output_path.exists():
            self._restore()
        elif not self.setup_error:
            self.plan = make_plan(len(self.students), self.universities)

    def _metadata(self, cursor: int) -> dict:
        return {"roster": fingerprint(self.students), "universities": self.universities,
                "plan": self.plan, "cursor": cursor}

    def _results(self, cursor: int | None = None) -> list[list[str | None]]:
        cursor = self.cursor if cursor is None else cursor
        return [[self.plan[i][r] if 2 * i + r < cursor else None for r in range(2)]
                for i in range(len(self.students))]

    def _restore(self):
        """恢复名单与规则一致、未被改写的结果，拒绝覆盖不匹配的文件。"""
        try:
            wb = load_workbook(self.output_path, read_only=True, data_only=False)
            try:
                description = wb.properties.description or ""
                if not description.startswith(METADATA_PREFIX):
                    raise ValueError("不是本程序生成的结果文件")
                metadata = json.loads(description[len(METADATA_PREFIX):])
                if metadata["roster"] != fingerprint(self.students):
                    raise ValueError("学生名单或姓名顺序已变化")
                if metadata["universities"] != self.universities:
                    raise ValueError("大学列表已变化")
                plan, cursor = metadata["plan"], metadata["cursor"]
                if type(cursor) is not int or not 0 <= cursor <= 2 * len(self.students):
                    raise ValueError("抽签进度无效")
                if (not isinstance(plan, list) or len(plan) != len(self.students)
                        or any(not isinstance(pair, list) or len(pair) != 2
                               or pair[0] == pair[1]
                               or any(u not in self.universities for u in pair) for pair in plan)
                        or any(len({pair[r] for pair in plan}) != len(plan) for r in range(2))):
                    raise ValueError("结果不符合两场分配规则")
                self.plan = plan
                expected = [tuple(HEADERS)] + [
                    (i + 1, name, *pair)
                    for i, (name, pair) in enumerate(zip(self.students, self._results(cursor)))
                ]
                if list(wb.worksheets[0].iter_rows(values_only=True)) != expected:
                    raise ValueError("结果表中的姓名或抽签内容已被修改")
                self.cursor, self.restored = cursor, True
            finally:
                wb.close()
        except Exception as exc:
            raise ValueError(
                f"无法续接 {self.output_path.name}：{exc}。"
                "请先将旧结果移走或改名备份，再重新运行；程序不会覆盖不匹配的结果。"
            ) from exc

    def snapshot(self) -> dict:
        with self.lock:
            results = self._results()
            pools = [[u for u in self.universities if u not in {p[r] for p in results}]
                     for r in range(2)]
            done = self.cursor == len(self.students) * 2 and not self.setup_error
            available = pools[self.cursor % 2] if not done else []
            if not done and self.cursor % 2:
                available = [u for u in available if u != results[self.cursor // 2][0]]
            return {
                "students": self.students, "universities": self.universities,
                "results": results, "cursor": self.cursor,
                "total_draws": len(self.students) * 2,
                "completed_students": self.cursor // 2, "done": done,
                "session_id": self.session_id, "restored": self.restored,
                "output_name": self.output_path.name, "rule_text": RULE_TEXT,
                "setup_error": self.setup_error, "available_by_round": pools,
                "available_universities": available,
            }

    def _write(self, cursor: int):
        """临时文件与原子替换；失败时保持旧文件和原有抽签进度。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "大学节抽签结果"
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "C2"
        ws.append(HEADERS)
        for i, (name, pair) in enumerate(zip(self.students, self._results(cursor)), 1):
            ws.append([i, name, *pair])
            ws.cell(i + 1, 2).data_type = "s"
        for row in ws:
            for cell in row:
                cell.font = Font(name="Arial", size=11, color="243247")
                cell.alignment = Alignment(vertical="center", horizontal="right" if cell.column == 1 else "left")
            ws.row_dimensions[row[0].row].height = 27
            if row[0].row > 1 and row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor="F0F4F8")
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="17263F")
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col, width in {"A": 9, "B": 22, "C": 36, "D": 36}.items():
            ws.column_dimensions[col].width = width
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        wb.properties.title = "大学节讲座分配"
        wb.properties.description = METADATA_PREFIX + json.dumps(self._metadata(cursor), ensure_ascii=False)
        temp_path = None
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(dir=self.output_path.parent, prefix=".lottery-",
                                             suffix=".xlsx", delete=False) as handle:
                temp_path = Path(handle.name)
            wb.save(temp_path)
            with temp_path.open("r+b") as handle:
                os.fsync(handle.fileno())
            os.replace(temp_path, self.output_path)
        finally:
            wb.close()
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)

    def draw(self, session_id: str, cursor: int) -> dict:
        with self.lock:
            if session_id != self.session_id or type(cursor) is not int or cursor != self.cursor:
                raise ConflictError("页面进度已变化，请同步最新结果后再抽取。")
            if self.setup_error:
                raise ValueError(self.setup_error)
            if self.cursor >= len(self.students) * 2:
                raise ConflictError("全部抽签已完成。")
            student_index, round_index = divmod(self.cursor, 2)
            self._write(self.cursor + 1)
            self.cursor += 1
            return {"state": self.snapshot(), "draw": {
                "student_index": student_index, "round_index": round_index,
                "university": self.plan[student_index][round_index],
            }}


class OutputLock:
    """防止多个程序实例同时写结果；进程退出后操作系统自动释放锁。"""
    def __init__(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.handle = path.with_name("." + path.name + ".lock").open("a+b")
        try:
            if os.name == "nt":
                import msvcrt
                self.handle.seek(0, 2)
                if not self.handle.tell():
                    self.handle.write(b"0")
                    self.handle.flush()
                self.handle.seek(0)
                msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            self.handle.close()
            raise ValueError("已有程序正在使用此结果文件，请先关闭原来的程序。") from exc

    def close(self):
        self.handle.close()


def make_handler(lottery: Lottery):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_):
            pass

        def send_bytes(self, data: bytes, content_type: str, status: int = 200,
                       download: bool = False):
            self.send_response(status)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            if download:
                self.send_header("Content-Disposition", 'attachment; filename="result.xlsx"')
            self.end_headers()
            try:
                self.wfile.write(data)
            except (BrokenPipeError, ConnectionResetError):
                pass

        def send_json(self, payload, status=200):
            self.send_bytes(json.dumps(payload, ensure_ascii=False).encode("utf-8"),
                            "application/json; charset=utf-8", status)

        def valid_host(self):
            return self.headers.get("Host") == f"127.0.0.1:{self.server.server_port}"

        def do_GET(self):
            if not self.valid_host():
                return self.send_json({"error": "请使用程序显示的本地地址。"}, 403)
            path = urlsplit(self.path).path
            if path == "/":
                self.send_bytes((ASSET_DIR / "lottery.html").read_bytes(), "text/html; charset=utf-8")
            elif path == "/api/state":
                self.send_json(lottery.snapshot())
            elif path == "/api/export":
                with lottery.lock:
                    if not lottery.cursor:
                        return self.send_json({"error": "尚未抽取，没有可导出的结果。"}, 400)
                    try:
                        data = lottery.output_path.read_bytes()
                    except OSError:
                        return self.send_json({"error": "无法读取结果文件，请检查文件是否被移走。"}, 500)
                    self.send_bytes(data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", download=True)
            else:
                self.send_error(404)

        def do_POST(self):
            origin = self.headers.get("Origin")
            if not self.valid_host() or (origin and origin != f"http://{self.headers['Host']}"):
                return self.send_json({"error": "只允许当前本地页面操作。"}, 403)
            if urlsplit(self.path).path != "/api/draw":
                return self.send_json({"error": "接口不存在。"}, 404)
            try:
                size = int(self.headers.get("Content-Length", 0))
                if not 0 < size <= 4096:
                    raise ValueError("请求大小无效。")
                payload = json.loads(self.rfile.read(size))
                if not isinstance(payload, dict):
                    raise ValueError("请求格式无效。")
                result = lottery.draw(payload.get("session_id"), payload.get("cursor"))
                self.send_json(result)
            except ConflictError as exc:
                self.send_json({"error": str(exc)}, 409)
            except (ValueError, UnicodeError) as exc:
                self.send_json({"error": str(exc)}, 400)
            except OSError:
                self.send_json({"error": "结果保存失败，本次未计入抽签。请关闭 Excel 中的结果文件，检查文件夹写入权限后重试。"}, 500)
            except Exception:
                self.send_json({"error": "本次抽签未完成，请重新加载页面确认已保存进度。"}, 500)
    return Handler


def configure_console_output():
    """重定向到 Windows 日志时也使用 UTF-8，避免中文触发 cp1252 编码错误。

    PyInstaller 的程序可能忽略 PYTHONIOENCODING，因此必须在程序内设置，
    并且要早于 argparse 的 --help 和错误信息输出。
    """
    for stream in (sys.stdout, sys.stderr):
        if stream is not None and hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="backslashreplace")


def main():
    configure_console_output()
    parser = argparse.ArgumentParser(description="大学节 SVG 抽签：同场不重复，每人两场去不同大学。")
    parser.add_argument("--input", type=Path, default=INPUT_FILE, help="学生名单，默认程序旁的 student.xlsx")
    parser.add_argument("--output", type=Path, default=OUTPUT_FILE, help="结果文件，默认程序旁的 result.xlsx")
    parser.add_argument("--port", type=int, default=0, help="本地端口，默认自动选择可用端口")
    parser.add_argument("--no-browser", action="store_true", help="不自动打开浏览器")
    args = parser.parse_args()
    output_lock = None
    server = None
    try:
        if args.input.resolve() == args.output.resolve():
            raise ValueError("名单与结果不能使用同一个文件。")
        if not 0 <= args.port <= 65535:
            raise ValueError("端口必须在 0–65535 之间。")
        if not (ASSET_DIR / "lottery.html").is_file():
            raise ValueError("缺少 lottery.html，请将它与 lottery.py 放在同一文件夹。")
        students = read_students(args.input)
        output_lock = OutputLock(args.output.resolve())
        lottery = Lottery(students, args.output.resolve())
        server = ThreadingHTTPServer(("127.0.0.1", args.port), make_handler(lottery))
        server.daemon_threads = True
        url = f"http://127.0.0.1:{server.server_port}"
        print(f"大学节抽签已启动：{url}\n已读取 {len(students)} 位同学，{len(UNIVERSITIES)} 所大学。", flush=True)
        print(f"{RULE_TEXT}\n结果保存至：{lottery.output_path}\n关闭程序请按 Ctrl+C。", flush=True)
        if lottery.setup_error:
            print(lottery.setup_error, flush=True)
        if not args.no_browser:
            webbrowser.open(url)
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n程序已关闭，已抽出的结果保存在 Excel 中。")
    except (ValueError, OSError) as exc:
        print(f"无法启动：{exc}", file=sys.stderr)
        if getattr(sys, "frozen", False) and sys.stdin and sys.stdin.isatty():
            input("按回车关闭窗口……")
        return 1
    finally:
        if server is not None:
            server.server_close()
        if output_lock is not None:
            output_lock.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
