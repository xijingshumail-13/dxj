"""Check a bundled EXE from another directory, using synthetic Excel files.

Usage: python tests/smoke_packaged.py dist/university-lottery.exe
A .py path is also accepted for a local pre-packaging check.
"""
from __future__ import annotations

import argparse
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import time
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook


def request(base, path, payload=None):
    data = None if payload is None else json.dumps(payload).encode("utf-8")
    req = Request(base + path, data=data, headers={"Content-Type": "application/json", "Origin": base})
    with urlopen(req, timeout=5) as response:
        return response.read()


def stop_process(process):
    if process is None:
        return
    if os.name == "nt":
        # PyInstaller onefile has a bootloader parent and a Python child.
        subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
    elif process.poll() is None:
        process.terminate()
    try:
        process.wait(timeout=10)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=10)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("program", type=Path)
    args = parser.parse_args()
    source = args.program.resolve(strict=True)
    with tempfile.TemporaryDirectory(prefix="lottery-smoke-") as temporary:
        root = Path(temporary)
        application = root / "抽签程序"
        other_directory = root / "different-working-directory"
        application.mkdir()
        other_directory.mkdir()
        target = application / source.name
        shutil.copy2(source, target)
        if source.suffix.lower() == ".py":
            shutil.copy2(source.with_name("lottery.html"), application / "lottery.html")
            command = [sys.executable, str(target)]
        else:
            command = [str(target)]
        command += ["--no-browser", "--port", "0"]
        wb = Workbook()
        wb.active.append(["姓名"])
        wb.active.append(["集成测试甲"])
        wb.active.append(["集成测试乙"])
        wb.save(application / "student.xlsx")
        wb.close()
        log = root / "program.log"
        process = None
        try:
            with log.open("wb") as output:
                process = subprocess.Popen(command, cwd=other_directory,
                                           stdin=subprocess.DEVNULL, stdout=output, stderr=subprocess.STDOUT)
                deadline = time.monotonic() + 60
                base = None
                while time.monotonic() < deadline:
                    if process.poll() is not None:
                        raise AssertionError(f"Program exited early ({process.returncode}).")
                    match = re.search(rb"http://127\.0\.0\.1:\d+", log.read_bytes())
                    if match:
                        base = match.group().decode("ascii")
                        try:
                            state = json.loads(request(base, "/api/state"))
                            break
                        except (URLError, HTTPError):
                            pass
                    time.sleep(0.1)
                else:
                    raise AssertionError("Program did not start within 60 seconds.")
                assert base is not None
                page = request(base, "/").decode("utf-8")
                assert '<svg id="slotMachine"' in page, "Bundled SVG page is missing."
                assert state["students"] == ["集成测试甲", "集成测试乙"]
                assert state["cursor"] == 0 and not state["setup_error"]
                for cursor in range(4):
                    reply = json.loads(request(base, "/api/draw", {"session_id": state["session_id"], "cursor": cursor}))
                    state = reply["state"]
                    assert state["cursor"] == cursor + 1
                    assert (application / "result.xlsx").is_file(), "Draw did not save beside executable."
                assert state["done"]
                assert all(a != b for a, b in state["results"])
                assert all(len({pair[r] for pair in state["results"]}) == 2 for r in (0, 1))
                saved = (application / "result.xlsx").read_bytes()
                exported = request(base, "/api/export")
                assert exported == saved, "Export differs from the saved workbook."
                result = load_workbook(io.BytesIO(exported), read_only=True, data_only=False)
                try:
                    rows = list(result.active.iter_rows(values_only=True))
                    assert rows[1:] == [(i + 1, name, *state["results"][i]) for i, name in enumerate(state["students"])]
                finally:
                    result.close()
                assert not (other_directory / "result.xlsx").exists(), "Output used working directory."
                print("Packaged smoke test passed: embedded SVG, executable-relative paths, all draws, save and export.")
        except Exception:
            print(log.read_bytes().decode("utf-8", errors="replace"), file=sys.stderr)
            raise
        finally:
            stop_process(process)


if __name__ == "__main__":
    main()
